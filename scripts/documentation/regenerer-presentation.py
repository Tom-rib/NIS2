#!/usr/bin/env python3
"""Regenere tous les formats de la presentation depuis la source HTML.

    python scripts/documentation/regenerer-presentation.py

Produit, depuis `Documentation/presentation-soutenance.html` :
  - Documentation/presentation-soutenance.pdf    projection et impression
  - Documentation/slides-png/slide-NN.png        2560x1440, une image par slide
  - Documentation/slides-png.zip                 les memes PNG en archive
  - Documentation/presentation-soutenance.pptx   import Canva / PowerPoint / Slides

Pourquoi le PPTX plutot que le PDF pour Canva : un PDF ne contient pas de tableaux,
seulement du texte positionne et des traits. Canva recree une zone de texte par
fragment et recalcule la mise en page avec ses propres polices ; Segoe UI lui
manquant, les metriques changent et les blocs glissent. Une image plein cadre par
slide supprime le probleme.

Pourquoi Playwright et non `msedge --screenshot` : quand une instance Edge tourne
deja (cas courant), la ligne de commande lui est deleguee et les options headless
sont ignorees en silence -- code retour 0, aucun fichier ecrit. Playwright lance
son propre processus et n'a pas ce probleme.

Prerequis :
    pip install playwright python-pptx
Playwright utilise le Edge deja installe (channel msedge), sans telechargement.
"""

from __future__ import annotations

import pathlib
import shutil
import sys
import zipfile

RACINE = pathlib.Path(__file__).resolve().parents[2]
SOURCE = RACINE / "Documentation" / "presentation-soutenance.html"
PDF = RACINE / "Documentation" / "presentation-soutenance.pdf"
PPTX = RACINE / "Documentation" / "presentation-soutenance.pptx"
DOSSIER_PNG = RACINE / "Documentation" / "slides-png"
ARCHIVE = RACINE / "Documentation" / "slides-png.zip"

LARGEUR, HAUTEUR = 1280, 720
ECHELLE = 2  # 1280x720 * 2 = 2560x1440


def main() -> None:
    if not SOURCE.exists():
        sys.exit(f"Source introuvable : {SOURCE}")

    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        sys.exit("playwright est absent.  pip install playwright")
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        sys.exit("python-pptx est absent.  pip install python-pptx")

    if DOSSIER_PNG.exists():
        shutil.rmtree(DOSSIER_PNG)
    DOSSIER_PNG.mkdir(parents=True)

    with sync_playwright() as pw:
        try:
            nav = pw.chromium.launch(channel="msedge")
        except Exception:
            nav = pw.chromium.launch()  # chromium embarque, si Edge est absent

        # --- 1. PDF : media "print", la taille vient du @page de la feuille ---
        page = nav.new_page()
        page.goto(SOURCE.as_uri(), wait_until="networkidle")
        nb = page.locator("section.slide").count()
        if nb == 0:
            sys.exit("Aucune slide trouvee dans la source.")
        print(f"Source : {SOURCE.name} - {nb} slides")

        page.pdf(path=str(PDF), prefer_css_page_size=True, print_background=True)
        print(f"  PDF   : {PDF.name} ({PDF.stat().st_size // 1024} Ko)")
        page.close()

        # --- 2. Un PNG par slide, capture element par element ---------------
        vue = nav.new_page(
            viewport={"width": LARGEUR, "height": HAUTEUR},
            device_scale_factor=ECHELLE,
        )
        vue.goto(SOURCE.as_uri(), wait_until="networkidle")
        vue.emulate_media(media="screen")
        # marges et ombres du mode ecran retirees : la capture doit etre a ras bord
        vue.add_style_tag(content="body{background:#fff;margin:0}"
                                  ".slide{margin:0!important;box-shadow:none!important}")

        slides = vue.locator("section.slide")
        for i in range(nb):
            cible = DOSSIER_PNG / f"slide-{i + 1:02d}.png"
            slides.nth(i).screenshot(path=str(cible))
            print(f"  PNG   : {cible.name}", end="\r")
        vue.close()
        nav.close()

    images = sorted(DOSSIER_PNG.glob("slide-*.png"))
    if len(images) != nb:
        sys.exit(f"Rendu incomplet : {len(images)} images pour {nb} slides")

    import struct
    l, h = struct.unpack(">II", images[0].read_bytes()[16:24])
    print(f"  PNG   : {len(images)} images en {l}x{h}        ")

    # --- 3. Archive ZIP ------------------------------------------------------
    with zipfile.ZipFile(ARCHIVE, "w", zipfile.ZIP_DEFLATED) as z:
        for img in images:
            z.write(img, img.name)
    print(f"  ZIP   : {ARCHIVE.name} ({ARCHIVE.stat().st_size // 1024 // 1024} Mo)")

    # --- 4. PPTX : une image plein cadre par slide ---------------------------
    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 13,333 pouces
    prs.slide_height = Emu(6858000)   #  7,5   pouces -> 16:9 exact
    vierge = prs.slide_layouts[6]
    for img in images:
        s = prs.slides.add_slide(vierge)
        s.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(PPTX)
    print(f"  PPTX  : {PPTX.name} ({PPTX.stat().st_size // 1024 // 1024} Mo, {len(prs.slides)} slides)")

    print("\nLes quatre formats sont a jour.")


if __name__ == "__main__":
    main()
